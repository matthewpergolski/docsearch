"""Build the tiny end-to-end fixture used by scripts/verify.sh.

Writes tests/fixtures/verify/files/ (handbook.docx, report.pdf with 2 pages, budget.xlsx with 2 sheets,
suppliers.csv), tests/fixtures/verify/questions.jsonl and tests/fixtures/verify/meta.expected.jsonl. The content is
invented: a fictional small bakery. Each question's answer phrase occurs in exactly one file (and, for the PDF, on
exactly one page), so every question has one clear gold unit. The script checks that rule, and checks the PDF page
split after conversion.

Each Office file and the PDF carry a title, an author and a date of their own in the file's properties; the dates
differ by file so a date filter can leave one out. meta.expected.jsonl holds the rows `docsearch meta` and `sync`
must write for them.

Run: uv run python scripts/make_verify_fixture.py
Needs LibreOffice (`soffice`) for the PDF.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "verify"
FILES = OUT / "files"
SOURCE = "verify"
FIXED_TIME = datetime(2026, 1, 5, 9, 0, 0)
AUTHOR = "verify fixture"
HANDBOOK_DATE = datetime(2024, 3, 4, 9, 0, 0)
REPORT_DATE = datetime(2025, 9, 30, 9, 0, 0)
BUDGET_TITLE = "Brindlewick Bakery budget 2026"

HANDBOOK_TITLE = "Brindlewick Bakery staff handbook"
HANDBOOK = [
    "Welcome to the team. This handbook explains how the bakery runs from day to day.",
    "Shifts. The early shift starts at 04:30 and ends at 12:30. Bakers sign in at the back door.",
    "Paid leave. Every new baker receives 18 days of paid leave each year, plus public holidays. "
    "Ask the shift lead to book leave at least two weeks ahead.",
    "Uniform. Aprons and caps are washed by an outside laundry. Put used aprons in the blue bin.",
    "Safety. Report any burn or cut to the shift lead on the same day.",
]

REPORT_TITLE = "Brindlewick Bakery summer report"
REPORT_PAGES = [
    [
        "Bread sales",
        "The rye sourdough line sold 4,200 loaves in July. The new stone mill gives a coarser flour, "
        "and regular customers now ask for the rye by name.",
        "Two local restaurants now take a daily order of seeded rolls.",
    ],
    [
        "Customer survey",
        "In the August survey, 71 percent of café visitors asked for longer weekend opening hours. "
        "From October the café stays open until 20:00 on Saturdays.",
        "The team will repeat the survey in the spring.",
    ],
]

BUDGET = {
    "Equipment": [
        ("Item", "Cost (EUR)", "Quarter"),
        ("Deck oven replacement", 12500, "Q1"),
        ("Dough sheeter", 3800, "Q2"),
        ("Walk-in cooler repair", 2150, "Q3"),
    ],
    "Marketing": [
        ("Item", "Cost (EUR)", "Quarter"),
        ("Farmers market stall fee", 640, "Q2"),
        ("Printed menu cards", 310, "Q1"),
        ("Local radio advert", 1900, "Q4"),
    ],
}

SUPPLIERS = [
    ("Supplier", "Product", "Delivery day"),
    ("Quillfeather Grain Co-op", "Organic spelt flour", "Monday"),
    ("Larkspur Hollow Dairy", "Cultured butter", "Wednesday"),
    ("Ottermoor Orchards", "Cooking apples", "Friday"),
]


@dataclass(frozen=True)
class Fact:
    qid: str
    query: str
    file: str
    page: int | None  # PDF questions are judged per page; the other formats per file
    answer: str  # the phrase that must occur in the gold unit and in no other unit

    @property
    def unit(self) -> str:
        return f"{SOURCE}/{self.file}" + (f"#p{self.page}" if self.page else "")


FACTS = [
    Fact("verify-01", "How many days of paid leave does a new baker get each year?", "handbook.docx", None,
         "18 days of paid leave"),
    Fact("verify-02", "How many rye sourdough loaves were sold in July?", "report.pdf", 1, "4,200 loaves"),
    Fact("verify-03", "What share of café visitors asked for longer weekend opening hours?", "report.pdf", 2,
         "71 percent"),
    Fact("verify-04", "What is the budgeted cost of the deck oven replacement?", "budget.xlsx", None,
         "Deck oven replacement"),
    Fact("verify-05", "How much does the farmers market stall fee cost?", "budget.xlsx", None,
         "Farmers market stall fee"),
    Fact("verify-06", "Which day does the cultured butter supplier deliver?", "suppliers.csv", None,
         "Cultured butter"),
]  # fmt: skip


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def unit_texts() -> dict[str, str]:
    """The text of every gold-able unit, from the same tables the files are written from."""
    rows = lambda tbl: " ".join(str(c) for r in tbl for c in r)  # noqa: E731
    texts = {f"{SOURCE}/handbook.docx": " ".join([HANDBOOK_TITLE, *HANDBOOK])}
    for i, page in enumerate(REPORT_PAGES, 1):
        texts[f"{SOURCE}/report.pdf#p{i}"] = " ".join(([REPORT_TITLE] if i == 1 else []) + page)
    texts[f"{SOURCE}/budget.xlsx"] = " ".join(f"{name} {rows(tbl)}" for name, tbl in BUDGET.items())
    texts[f"{SOURCE}/suppliers.csv"] = rows(SUPPLIERS)
    return texts


def check_unique(texts: dict[str, str]) -> None:
    for f in FACTS:
        holders = [u for u, t in texts.items() if _norm(f.answer) in _norm(t)]
        if holders != [f.unit]:
            raise SystemExit(f"{f.qid}: answer {f.answer!r} must occur only in {f.unit}, found in {holders}")


def write_handbook(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.core_properties.title = HANDBOOK_TITLE
    doc.core_properties.author = AUTHOR
    doc.core_properties.created = doc.core_properties.modified = HANDBOOK_DATE
    doc.add_heading(HANDBOOK_TITLE, level=1)
    for p in HANDBOOK:
        doc.add_paragraph(p)
    doc.save(path)


def find_soffice() -> str:
    for cand in (shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if cand and Path(cand).exists():
            return cand
    raise SystemExit("LibreOffice (soffice) not found; it is needed to write report.pdf")


def write_report(path: Path) -> None:
    from docx import Document

    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        doc = Document()
        doc.core_properties.title = REPORT_TITLE
        doc.core_properties.author = AUTHOR
        doc.core_properties.created = doc.core_properties.modified = REPORT_DATE
        doc.add_heading(REPORT_TITLE, level=1)
        for i, page in enumerate(REPORT_PAGES):
            if i:
                doc.add_page_break()
            doc.add_heading(page[0], level=2)
            for p in page[1:]:
                doc.add_paragraph(p)
        src = tmp / "report.docx"
        doc.save(src)
        # A private profile: the conversion must not depend on (or fail against) a LibreOffice window already open.
        profile = f"-env:UserInstallation={(tmp / 'lo-profile').as_uri()}"
        subprocess.run([find_soffice(), profile, "--headless", "--norestore", "--convert-to", "pdf",
                        "--outdir", str(tmp), str(src)], check=True, capture_output=True, timeout=180)  # fmt: skip
        shutil.move(tmp / "report.pdf", path)
    set_pdf_creation_date(path, REPORT_DATE)


def set_pdf_creation_date(path: Path, when: datetime) -> None:
    """LibreOffice keeps the Word title and author but stamps the conversion time as the PDF CreationDate. Write
    the report's own date in its place, at the same length, so the cross-reference offsets stay valid."""
    data = path.read_bytes()
    hits = list(re.finditer(rb"/CreationDate\((D:[^)]*)\)", data))
    if len(hits) != 1:
        raise SystemExit(f"report.pdf: expected one /CreationDate string, found {len(hits)}")
    old = hits[0].group(1)
    stamp = when.strftime("D:%Y%m%d%H%M%S").encode()
    new = next((stamp + tz for tz in (b"+00'00'", b"+00'00", b"Z", b"") if len(stamp + tz) == len(old)), None)
    if new is None:
        raise SystemExit(f"report.pdf: CreationDate {old!r} has an unexpected length")
    path.write_bytes(data[: hits[0].start(1)] + new + data[hits[0].end(1) :])


def check_pdf_pages(path: Path) -> None:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(path))
    try:
        pages = [_norm(doc[i].get_textpage().get_text_range()) for i in range(len(doc))]
        info = doc.get_metadata_dict()
    finally:
        doc.close()
    want = (REPORT_TITLE, AUTHOR, REPORT_DATE.strftime("D:%Y%m%d"))
    got = (info.get("Title"), info.get("Author"), info.get("CreationDate", "")[:10])
    if got != want:
        raise SystemExit(f"report.pdf metadata (title, author, creation day) is {got}, expected {want}")
    if len(pages) != len(REPORT_PAGES):
        raise SystemExit(f"report.pdf has {len(pages)} pages, expected {len(REPORT_PAGES)}")
    for f in FACTS:
        if f.file == "report.pdf":
            on = [i for i, t in enumerate(pages, 1) if _norm(f.answer) in t]
            if on != [f.page]:
                raise SystemExit(f"{f.qid}: {f.answer!r} is on PDF page(s) {on}, expected [{f.page}]")


def write_budget(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.title = BUDGET_TITLE
    wb.properties.creator = AUTHOR
    wb.properties.created = FIXED_TIME
    wb.remove(wb.active)
    for name, rows in BUDGET.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    wb.save(path)
    set_xlsx_modified(path, FIXED_TIME)


def set_xlsx_modified(path: Path, when: datetime) -> None:
    """openpyxl stamps the save time as the modified date; write the fixed one into docProps/core.xml."""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    stamp = when.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                data, n = re.subn(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, data)
                if n != 1:
                    raise SystemExit("budget.xlsx: no modified date in docProps/core.xml")
            z.writestr(info, data)


def write_suppliers(path: Path) -> None:
    with path.open("w", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows(SUPPLIERS)


def write_questions(path: Path) -> None:
    fmt = {".docx": "docx", ".pdf": "pdf", ".xlsx": "xlsx", ".csv": "csv"}
    with path.open("w") as out:
        for f in FACTS:
            row = {"qid": f.qid, "query": f.query, "source": SOURCE, "format": fmt[Path(f.file).suffix],
                   "level": "page" if f.page else "file", "gold": {f.unit: 1}, "meta": {"answer": f.answer}}  # fmt: skip
            out.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_expected_meta(path: Path) -> None:
    """The meta.jsonl rows for the fixture files once they sit at the top of the corpus folder (as in verify.sh)."""
    day = lambda dt: dt.date().isoformat()  # noqa: E731
    rows = [
        ("budget.xlsx", "xlsx", BUDGET_TITLE, AUTHOR, day(FIXED_TIME), day(FIXED_TIME), "file properties"),
        ("handbook.docx", "docx", HANDBOOK_TITLE, AUTHOR, day(HANDBOOK_DATE), day(HANDBOOK_DATE), "file properties"),
        ("report.pdf", "pdf", REPORT_TITLE, AUTHOR, day(REPORT_DATE), "", "pdf metadata"),
        ("suppliers.csv", "csv", "", "", "", "", "none"),
    ]
    with path.open("w") as out:
        for name, fmt, title, author, created, modified, source in rows:
            row = {"doc_id": f"{SOURCE}/{name}", "path": name, "format": fmt, "title": title, "author": author,
                   "created": created, "modified": modified, "date_source": source, "folder": "", "file_name": name}  # fmt: skip
            out.write(json.dumps(row, ensure_ascii=False) + "\n")


def main() -> int:
    check_unique(unit_texts())
    FILES.mkdir(parents=True, exist_ok=True)
    write_handbook(FILES / "handbook.docx")
    write_report(FILES / "report.pdf")
    check_pdf_pages(FILES / "report.pdf")
    write_budget(FILES / "budget.xlsx")
    write_suppliers(FILES / "suppliers.csv")
    write_questions(OUT / "questions.jsonl")
    write_expected_meta(OUT / "meta.expected.jsonl")
    for p in sorted(FILES.iterdir()):
        print(f"wrote {p.relative_to(OUT.parents[2])} ({p.stat().st_size} bytes)")
    print(f"wrote {(OUT / 'questions.jsonl').relative_to(OUT.parents[2])} ({len(FACTS)} questions)")
    print(f"wrote {(OUT / 'meta.expected.jsonl').relative_to(OUT.parents[2])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
