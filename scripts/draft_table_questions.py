"""Draft the held-out table question set (Phase C3) from the original files with the Grok CLI.

  uv run python scripts/draft_table_questions.py --plan           targets, look-alike pairs, split; no LLM call
  uv run python scripts/draft_table_questions.py --split dev --limit 20   pilot -> work/cache/table_questions/pilot.jsonl
  uv run python scripts/draft_table_questions.py                  all targets -> data/questions_tables.jsonl

Source of truth is the original file: PDF page text (pypdfium2), xlsx sheets (openpyxl, data_only, hidden sheets
skipped), csv rows, docx and pptx tables. Gold units follow the corpus: "<doc_id>#p<page>" (1-based, equal to
Docling's page_no) for pdf and pptx; the file for docx (Docling gives docx no pages), xlsx and csv.
Every Grok reply is cached by prompt hash under work/cache/table_questions/, so a rerun resumes where it stopped.
The dev/heldout split is by document family and is frozen in split.json on the first run.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import random
import re
import statistics
import sys
import threading
import time
import warnings
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docsearch.config import load_config
from docsearch.draft_questions import extract_json
from docsearch.llm import LLMClient, LLMError, settings
from docsearch.questions import ngram_overlap

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "corpus"
CACHE = ROOT / "work" / "cache" / "table_questions"
OUT = ROOT / "data" / "questions_tables.jsonl"
FORMATS = ("pdf", "xlsx", "csv", "docx", "pptx")
PAGE_LEVEL = {"pdf", "pptx"}
CAPS = {"pdf": (40, 3), "xlsx": (8, 2), "pptx": (8, 3), "docx": (0, 2), "csv": (0, 1)}  # (look-alike, single) per file
MIN_LABELS = {"pdf": 5, "xlsx": 4, "csv": 3, "docx": 3, "pptx": 3}
HELDOUT_SHARE = 0.3
PAIR_JACCARD, VALUE_JACCARD_MAX = 0.8, 0.5
MAX_CHARS = 6000
MAX_ROWS = 300
COPY_RUN = 6
PROBE_CALLS = 10
BUDGET_S = 6 * 3600
SEED = 31

NUM = re.compile(r"\(?[-−–]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\)?%?")
SPACED = re.compile(r"(?<![\d.,])\d{1,3}(?:[ \u00a0\u202f]\d{3})+(?![\d.,]?\d)")
CELL_NUM = re.compile(r"[$€£]?\(?[-−–+]?[$€£]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\)?%?")
NAMES_PLACE = re.compile(
    r"\b(page|slide|table|figure|annex|chart|row|column)\s+[a-z]?\d|\b(sheet|tab|file|spreadsheet|workbook|"
    r"document|excerpt)\b",
    re.I,
)

DRAFT = """You write evaluation questions for a document search engine over a large collection of reports,
spreadsheets and slides. Below is the text of ONE {what} from a {fmt} file, taken from the original file.
It holds a table of numbers.{grid_note}

Pick ONE data cell that holds a number (not a header, not a year label, not a footnote mark) and write ONE
question whose answer is exactly that cell's value.
Rules:
- Word it as a user would type it: name the entity (row), the measure (what the table or column measures) and
  the period, e.g. "What was Croatia's employment rate for people aged 15-24 in 2023?".
- Give enough context (measure, group, unit base, period) that only this cell answers it.{pair_rule}
- Never put another number from the table into the question; years and age bands are fine.
- Never name the file, report, page, sheet, slide, table number or figure number.
- Use your own words: never copy a run of 5 or more words from the text.
- One sentence, at most 30 words, exactly one question mark.
- The answer is the cell value as printed, with its unit (e.g. "12.4%", "1,234 thousand", "EUR 3.2 billion").
Return JSON only, no prose:
{{"question": "...", "answer": "...", "value": "<the number exactly as printed in the cell>", "row_label": "...",
"column_header": "...", "cell": "<the cell ADDRESS if the text gives addresses, else empty>"}}

TARGET {what_upper}:
{text}
{partner}"""

PAIR_RULE = """
- A LOOK-ALIKE {what} follows the target. It has the same row labels but other values (another measure, period
  or group). Pick a cell whose value differs from the matching cell on the look-alike, and name the detail that
  tells the target apart, so the look-alike does not answer the question."""

REVIEW = """You review one evaluation question for a document search engine. A user types the question; the
search must find the TARGET {what}. Work only from the text below.
Keep the question only if ALL hold:
1. The TARGET text shows exactly this answer for what the question asks.
2. The question asks for one value and names entity, measure and period clearly enough that one cell answers it.
3. {partner_check}
4. The wording is natural; it does not name a file, report, page, sheet, slide, table or figure number, and it
   does not contain another number from the table (years and age bands are fine).
Return JSON only: {{"decision": "keep" or "drop", "reason": "<one short sentence>"}}

QUESTION: {question}
ANSWER: {answer}

TARGET {what_upper}:
{text}
{partner}"""

PARTNER_CHECK = "The LOOK-ALIKE text does not also answer the question (with the same or another value)."
NO_PARTNER_CHECK = "No other part of the TARGET text gives a different answer to the question."
PROMPT_VERSION = hashlib.sha256((DRAFT + PAIR_RULE + REVIEW).encode()).hexdigest()[:12]

log = logging.getLogger("draft_table_questions")


@dataclass
class Table:
    tid: str
    doc_id: str
    fmt: str
    source: str
    family: str
    loc: str
    unit: str
    text: str
    labels: frozenset[str]
    values: frozenset[Decimal]
    grid: dict[str, str] | None = None
    page: int | None = None
    sheet: str | None = None
    sheet_unit: str | None = None
    pair_id: str | None = None
    partner: Table | None = field(default=None, repr=False)

    @property
    def level(self) -> str:
        return "page" if self.fmt in PAGE_LEVEL else "file"

    @property
    def what(self) -> str:
        return {"pdf": "page", "pptx": "slide", "xlsx": "sheet", "csv": "table excerpt", "docx": "table"}[self.fmt]


def to_dec(tok: str) -> Decimal | None:
    t = tok.strip()
    neg = (t.startswith("(") and t.rstrip("%").endswith(")")) or t.lstrip("($€£")[:1] in "-−–"
    try:
        d = Decimal(re.sub(r"[^\d.]", "", t))
    except InvalidOperation:
        return None
    return -d if neg else d


def norm(d: Decimal) -> str:
    s = format(d.normalize(), "f")
    return "0" if s == "-0" else s


def join_spaced(text: str) -> str:
    return SPACED.sub(lambda m: re.sub(r"\s", "", m.group()), text)


def numbers(text: str) -> set[Decimal]:
    toks = NUM.findall(text) + NUM.findall(join_spaced(text))
    return {abs(d) for d in map(to_dec, toks) if d is not None}


def is_number(v: str) -> bool:
    return bool(CELL_NUM.fullmatch(v.replace(" ", "")))


def norm_label(s: str) -> str:
    return " ".join(re.sub(r"[^a-z ]", " ", s.lower()).split())


def jaccard(a: frozenset, b: frozenset) -> float:
    return len(a & b) / len(a | b) if a and b else 0.0


def col_letter(c: int) -> str:
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def family_of(doc_id: str) -> str:
    stem = Path(doc_id.split("/", 1)[1]).stem.lower()
    stem = re.sub(r"-k[a-z]\w{6,}$|-t[a-z]\w{6,}$", "", stem)
    stem = re.sub(r"[-_ ]?(edited|improved|small|modified|final)\b|[-_ ]?\(\d+\)|[-_ ]20\d\d$", "", stem)
    return stem.strip(" -_")


def clean_pdf(t: str) -> str:
    t = t.replace("\ufffe", "").replace("\u00ad", "").replace("\r\n", "\n")
    return "\n".join(line.strip() for line in t.split("\n") if line.strip())


def pdf_row(line: str) -> tuple[str, list[str]]:
    toks = line.split()
    i = 0
    while i < len(toks) and not NUM.fullmatch(toks[i]):
        i += 1
    return norm_label(" ".join(toks[:i])), [x for x in toks[i:] if NUM.fullmatch(x)]


def read_pdf(doc: dict) -> list[Table]:
    import pypdfium2 as pdfium

    out = []
    pdf = pdfium.PdfDocument(CORPUS / doc["path"])
    try:
        for i in range(len(pdf)):
            text = clean_pdf(pdf[i].get_textpage().get_text_range())[:MAX_CHARS]
            rows = [pdf_row(line) for line in text.split("\n")]
            labels = {lab for lab, nums in rows if len(nums) >= 3 and lab}
            vals = {abs(d) for _, nums in rows if len(nums) >= 3 for d in map(to_dec, nums) if d is not None}
            if len(labels) >= MIN_LABELS["pdf"]:
                out.append(_table(doc, f"p{i + 1}", f"{doc['doc_id']}#p{i + 1}", text, labels, vals, page=i + 1))
    finally:
        pdf.close()
    return out


def _table(doc: dict, loc: str, unit: str, text: str, labels, vals, **kw) -> Table:
    return Table(tid=f"{doc['doc_id']}@{loc}", doc_id=doc["doc_id"], fmt=doc["format"], source=doc["source"],
                 family=family_of(doc["doc_id"]), loc=loc, unit=unit, text=text, labels=frozenset(labels),
                 values=frozenset(vals), **kw)  # fmt: skip


def grid_table(doc: dict, loc: str, unit: str, rows: list[tuple[int, list[str]]], head: str, **kw) -> Table | None:
    grid, lines, labels, vals = {}, [head] if head else [], set(), set()
    for r, row in rows:
        cells = [(f"{col_letter(c)}{r}", v) for c, v in enumerate(row, 1) if v]
        if not cells:
            continue
        grid.update(cells)
        lines.append(f"{r}: " + " | ".join(f"{a}={v}" for a, v in cells))
        nums = [v for _, v in cells if is_number(v)]
        lab = next((norm_label(v) for _, v in cells if not is_number(v) and norm_label(v)), "")
        if len(nums) >= 1 and lab:
            labels.add(lab)
            vals.update(abs(d) for d in map(to_dec, nums) if d is not None)
    text = "\n".join(lines)[:MAX_CHARS]
    shown = {a: v for a, v in grid.items() if f"{a}={v}" in text}
    if len(labels) < MIN_LABELS[doc["format"]] or sum(map(is_number, shown.values())) < 6:
        return None
    return _table(doc, loc, unit, text, labels, vals, grid=shown, **kw)


def decimals(fmt: str) -> int | None:
    sec = fmt.split(";")[0]
    if "0" not in sec:
        return None
    m = re.search(r"\.(0+)", sec)
    return len(m.group(1)) if m else 0


def show(v, fmt: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        d = decimals(fmt or "")
        if "%" in (fmt or ""):
            return f"{v * 100:.{d or 0}f}%"
        if d is not None and isinstance(v, float):
            return f"{v:,.{d}f}" if "," in fmt else f"{v:.{d}f}"
        if isinstance(v, float):
            return norm(Decimal(str(round(v, 4))))
        return f"{v:,}" if "," in (fmt or "") else str(v)
    if isinstance(v, datetime):
        return v.date().isoformat() if not (v.hour or v.minute) else v.isoformat(sep=" ")
    return " ".join(str(v).split())


def read_xlsx(doc: dict) -> list[Table]:
    import openpyxl

    out = []
    wb = openpyxl.load_workbook(CORPUS / doc["path"], data_only=True, read_only=True)
    try:
        for si, ws in enumerate(wb.worksheets, 1):
            if ws.sheet_state != "visible":
                log.info("hidden sheet skipped: %s / %s", doc["doc_id"], ws.title)
                continue
            rows = []
            for r, row in enumerate(ws.iter_rows(), 1):
                if r > MAX_ROWS:
                    break
                rows.append((r, [show(c.value, getattr(c, "number_format", "")) for c in row[:30]]))
            t = grid_table(doc, ws.title, doc["doc_id"], rows, f"Sheet: {ws.title}", sheet=ws.title,
                           sheet_unit=f"{doc['doc_id']}#p{si}")  # fmt: skip
            if t:
                out.append(t)
    finally:
        wb.close()
    return out


def read_csv(doc: dict) -> list[Table]:
    with open(CORPUS / doc["path"], encoding="utf-8-sig", errors="replace", newline="") as f:
        rows = list(csv.reader(f))
    if len(rows) < 10 or len(rows[0]) < 3:
        log.info("csv skipped (under 3 columns or 10 rows): %s", doc["doc_id"])
        return []
    start = random.Random(f"{SEED}:{doc['doc_id']}").randrange(1, max(2, len(rows) - 40))
    window = [(1, rows[0][:20])] + [(r + 1, rows[r][:20]) for r in range(start, min(start + 40, len(rows)))]
    window = [(r, [" ".join(v.split()) for v in row]) for r, row in window]
    t = grid_table(doc, f"rows{start + 1}", doc["doc_id"], window, f"Header row 1, then rows {start + 1}-{start + 40}")
    return [t] if t else []


def read_docx(doc: dict) -> list[Table]:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable

    d = docx.Document(CORPUS / doc["path"])
    out, prev, k = [], [], 0
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            txt = " ".join("".join(t.text or "" for t in el.iter(qn("w:t"))).split())
            if txt:
                prev.append(txt)
        elif tag == "tbl":
            k += 1
            try:
                rows = [(r, [" ".join(c.text.split()) for c in row.cells]) for r, row in
                        enumerate(DocxTable(el, d).rows, 1)]  # fmt: skip
            except Exception as e:  # noqa: BLE001
                log.warning("docx table unreadable %s T%d: %s", doc["doc_id"], k, e)
                continue
            t = grid_table(
                doc, f"T{k}", doc["doc_id"], rows[:MAX_ROWS], "Text before the table: " + " / ".join(prev[-3:])
            )
            if t:
                out.append(t)
    return out


def read_pptx(doc: dict) -> list[Table]:
    import pptx

    out = []
    for si, slide in enumerate(pptx.Presentation(CORPUS / doc["path"]).slides, 1):
        texts = [" ".join(s.text_frame.text.split()) for s in slide.shapes if s.has_text_frame]
        tables = [s.table for s in slide.shapes if s.has_table]
        cands = []
        for tb in tables:
            rows = [(r, [" ".join(c.text.split()) for c in row.cells]) for r, row in enumerate(tb.rows, 1)]
            t = grid_table(doc, f"p{si}", f"{doc['doc_id']}#p{si}", rows,
                           "Slide text: " + " / ".join(x for x in texts if x), page=si)  # fmt: skip
            if t:
                cands.append(t)
        if cands:
            out.append(max(cands, key=lambda t: len(t.grid or {})))
    return out


READERS = {"pdf": read_pdf, "xlsx": read_xlsx, "csv": read_csv, "docx": read_docx, "pptx": read_pptx}


def load_docs() -> list[dict]:
    with open(CORPUS / "manifest.csv", newline="") as f:
        docs = [d for d in csv.DictReader(f) if d["format"] in FORMATS and (CORPUS / d["path"]).exists()]
    keep, seen = [], set()
    for d in sorted(docs, key=lambda d: (len(d["doc_id"]), d["doc_id"])):
        key = (d["format"], family_of(d["doc_id"]))
        if d["format"] not in PAGE_LEVEL and key in seen:
            log.info("variant skipped (same family, file-level gold would be ambiguous): %s", d["doc_id"])
            continue
        seen.add(key)
        keep.append(d)
    return sorted(keep, key=lambda d: d["doc_id"])


def pick(tables: list[Table], fmt: str) -> list[Table]:
    max_look, max_singles = CAPS[fmt]
    if max_look:
        for a in tables:
            best = max(((jaccard(a.labels, b.labels), b.tid, b) for b in tables
                        if b is not a and jaccard(a.values, b.values) <= VALUE_JACCARD_MAX), default=None)  # fmt: skip
            if best and best[0] >= PAIR_JACCARD:
                a.partner = best[2]
    look = [t for t in tables if t.partner]
    group = {t.tid: t.tid for t in look}

    def root(x: str) -> str:
        while group[x] != x:
            x = group[x]
        return x

    for t in look:
        group[root(t.tid)] = root(t.partner.tid)
    for t in look:
        t.pair_id = hashlib.sha256(root(t.tid).encode()).hexdigest()[:10]
    rest = [t for t in tables if not t.partner]
    rng = random.Random(f"{SEED}:{tables[0].doc_id}") if tables else random.Random(SEED)
    return rng.sample(look, min(max_look, len(look))) + rng.sample(rest, min(max_singles, len(rest)))


def build_targets() -> list[Table]:
    out = []
    for d in load_docs():
        try:
            tables = READERS[d["format"]](d)
        except Exception as e:  # noqa: BLE001
            log.warning("unreadable %s: %s", d["doc_id"], e)
            continue
        out += pick(tables, d["format"])
    return sorted(out, key=lambda t: t.tid)


def assign_split(targets: list[Table], path: Path) -> dict[str, str]:
    frozen = json.loads(path.read_text()) if path.exists() else {}
    split = dict(frozen)
    by_fmt: dict[str, Counter] = defaultdict(Counter)
    for t in targets:
        by_fmt[t.fmt][t.family] += 1
    for fams in by_fmt.values():
        goal = HELDOUT_SHARE * sum(fams.values())
        held = sum(n for f, n in fams.items() if split.get(f) == "heldout")
        for fam in sorted(
            (f for f in fams if f not in split), key=lambda f: hashlib.sha256(f"c3:{f}".encode()).hexdigest()
        ):
            take = abs(held + fams[fam] - goal) < abs(held - goal)
            split[fam] = "heldout" if take else "dev"
            held += fams[fam] if take else 0
    if split != frozen:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(dict(sorted(split.items())), indent=1) + "\n")
    return split


class Grok:
    def __init__(self, client: LLMClient):
        self.client = client
        self.latency: list[float] = []
        self.lock = threading.Lock()

    def __call__(self, prompt: str, kind: str) -> str:
        key = hashlib.sha256(prompt.encode()).hexdigest()
        path = CACHE / "calls" / f"{key}.json"
        if path.exists():
            return json.loads(path.read_text())["reply"]
        t0 = time.monotonic()
        reply = self.client.complete(prompt)
        dt = time.monotonic() - t0
        with self.lock:
            self.latency.append(dt)
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(f".{threading.get_ident()}.tmp")
        tmp.write_text(json.dumps({"kind": kind, "seconds": round(dt, 2), "reply": reply, "prompt": prompt}))
        tmp.replace(path)
        return reply


def partner_block(t: Table) -> str:
    return f"\nLOOK-ALIKE {t.what.upper()}:\n{t.partner.text}\n" if t.partner else ""


def grid_note(t: Table) -> str:
    return " Cells are written as ADDRESS=value (column letter + row number)." if t.grid else ""


def copied_words(q: str, t: Table) -> int:
    return ngram_overlap(q, [t.text] + ([t.partner.text] if t.partner else []))["longest"]


def value_keyed(q: str, t: Table, value: Decimal) -> bool:
    for d in numbers(q):
        if d == value or (d == d.to_integral_value() and (1900 <= d <= 2100 or d < 100)):
            continue
        if d in t.values:
            return True
    return False


def names_file(q: str, t: Table) -> bool:
    words = [w for w in re.split(r"[^a-z]+", Path(t.doc_id).stem.lower()) if len(w) >= 3]
    ql = " ".join(re.split(r"[^a-z]+", q.lower()))
    runs = [" ".join(words[i : i + 3]) for i in range(len(words) - 2)]
    return bool(NAMES_PLACE.search(q)) or any(r in ql for r in runs)


def locate_cell(t: Table, claimed: str, value: Decimal) -> tuple[str | None, str]:
    if not t.grid:
        return None, "page"
    a1 = (claimed or "").split("!")[-1].strip().upper()
    if a1 in t.grid and is_number(t.grid[a1]) and abs(to_dec(t.grid[a1]) or Decimal(-1)) == value:
        return a1, "claimed"
    hits = [a for a, v in t.grid.items() if is_number(v) and abs(to_dec(v) or Decimal(-1)) == value]
    return (hits[0], "unique_value") if len(hits) == 1 else (None, f"unverified ({len(hits)} cells hold the value)")


def draft_one(t: Table, qid: str, split: str, grok: Grok) -> dict:
    base = {"qid": qid, "tid": t.tid, "format": t.fmt, "split": split, "pair_id": t.pair_id}
    what = t.what
    prompt = DRAFT.format(what=what, what_upper=what.upper(), fmt=t.fmt, grid_note=grid_note(t), text=t.text,
                          pair_rule=PAIR_RULE.format(what=what) if t.partner else "", partner=partner_block(t))  # fmt: skip
    try:
        res = extract_json(grok(prompt, "draft"))
    except LLMError as e:
        return {**base, "reject": f"llm_error: {type(e).__name__}", "detail": str(e)[:200]}
    if not res or not res.get("question") or not (res.get("value") or res.get("answer")):
        return {**base, "reject": "bad_json"}
    q, answer = " ".join(str(res["question"]).split()), str(res.get("answer") or res.get("value"))
    toks = NUM.findall(join_spaced(str(res.get("value") or ""))) or NUM.findall(join_spaced(answer))
    value = to_dec(toks[0]) if toks else None
    draft = {**base, "question": q, "answer": answer, "draft": res}
    if value is None:
        return {**draft, "reject": "no_number_in_answer"}
    checks = [("value_not_on_page", abs(value) not in numbers(t.text)), ("value_in_question", abs(value) in numbers(q)),
              ("copied_6_words", copied_words(q, t) >= COPY_RUN), ("names_file_or_place", names_file(q, t)),
              ("over_one_question", q.count("?") != 1 or len(q.split()) > 40)]  # fmt: skip
    failed = [name for name, bad in checks if bad]
    if failed:
        return {**draft, "reject": failed[0], "failed": failed}
    a1, how = locate_cell(t, str(res.get("cell") or ""), abs(value))
    if t.grid and a1 is None:
        return {**draft, "reject": "cell_not_verified", "detail": how}
    review_prompt = REVIEW.format(what=what, what_upper=what.upper(), question=q, answer=answer, text=t.text,
                                  partner_check=PARTNER_CHECK if t.partner else NO_PARTNER_CHECK,
                                  partner=partner_block(t))  # fmt: skip
    try:
        rev = extract_json(grok(review_prompt, "review")) or {}
    except LLMError as e:
        return {**draft, "reject": f"review_llm_error: {type(e).__name__}", "detail": str(e)[:200]}
    decision = str(rev.get("decision", "")).lower()
    if decision not in ("keep", "drop"):
        return {**draft, "reject": "review_bad_json"}
    keyed = value_keyed(q, t, abs(value))
    kind = "table_valuekeyed" if keyed else "table_lookalike" if t.partner else "table_single"
    gold_cell = {"unit": t.sheet_unit or t.unit, "page": t.page, "sheet": t.sheet, "table": t.loc if t.fmt == "docx" else None,
                 "a1": a1, "a1_check": how, "row_label": res.get("row_label", ""), "column_header": res.get("column_header", "")}  # fmt: skip
    meta = {"type": kind, "status": decision, "backend": "cli", "provider": "grok", "model": grok.client.s.model,
            "answer": answer, "value_norm": norm(value), "gold_cell": {k: v for k, v in gold_cell.items() if v is not None},
            "pair_id": t.pair_id, "partner_unit": t.partner.unit if t.partner else None, "lookalike": bool(t.partner),
            "value_on_partner": bool(t.partner) and abs(value) in numbers(t.partner.text), "family": t.family,
            "split": split, "reviewer": "grok-cli", "review_reason": str(rev.get("reason", "")),
            "prompt_version": PROMPT_VERSION, "source_text": "original file"}  # fmt: skip
    row = {
        "qid": qid,
        "query": q,
        "source": t.source,
        "format": t.fmt,
        "level": t.level,
        "gold": {t.unit: 1},
        "meta": meta,
    }
    return {**draft, "row": row}


def pct(xs: list[float], p: float) -> float:
    xs = sorted(xs)
    return xs[min(len(xs) - 1, round(p * (len(xs) - 1)))]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows))
    tmp.replace(path)


def plan_report(targets: list[Table], split: dict[str, str]) -> None:
    c = Counter((t.fmt, "pair" if t.partner else "single", split[t.family]) for t in targets)
    print(f"targets {len(targets)}; prompt_version {PROMPT_VERSION}")
    for k, v in sorted(c.items()):
        print(f"  {k[0]:5s} {k[1]:6s} {k[2]:8s} {v}")
    fam = Counter((t.fmt, t.family, split[t.family]) for t in targets)
    for (fmt, f, s), n in sorted(fam.items()):
        print(f"  family {fmt:5s} {s:8s} {n:3d} {f}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--plan", action="store_true", help="print targets and split, call no LLM")
    ap.add_argument("--split", choices=("dev", "heldout", "all"), default="all")
    ap.add_argument("--limit", type=int, default=0, help="first N targets only (pilot)")
    ap.add_argument("--workers", type=int, default=4)
    a = ap.parse_args()
    CACHE.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        filename=CACHE / "run.log", level=logging.INFO, format="%(asctime)s %(name)s %(levelname)s %(message)s"
    )
    warnings.simplefilter("default")
    logging.captureWarnings(True)
    t_start = time.monotonic()
    targets = build_targets()
    split = assign_split(targets, CACHE / "split.json")
    qids = {t.tid: f"tables:{i}" for i, t in enumerate(targets)}
    plan_report(targets, split)
    if a.plan:
        return 0
    todo = sorted(
        (t for t in targets if a.split in ("all", split[t.family])),
        key=lambda t: hashlib.sha256(t.tid.encode()).hexdigest(),
    )
    todo = todo[: a.limit] if a.limit else todo
    full = a.split == "all" and not a.limit
    grok = Grok(LLMClient(settings(load_config(), "questions", backend="cli", provider="grok")))

    def work(t: Table) -> dict:
        return draft_one(t, qids[t.tid], split[t.family], grok)

    results: list[dict] = []
    with ThreadPoolExecutor(max(1, a.workers)) as ex:
        i = 0
        while i < len(todo) and len(grok.latency) < PROBE_CALLS:
            batch = todo[i : i + a.workers]
            results += list(ex.map(work, batch))
            i += len(batch)
        if grok.latency:
            med, p95 = statistics.median(grok.latency), pct(grok.latency, 0.95)
            calls_left = 2 * (len(todo) - i)
            proj = calls_left * med / max(1, a.workers)
            print(f"grok latency over first {len(grok.latency)} live calls: median {med:.1f}s, p95 {p95:.1f}s; "
                  f"projected {calls_left} calls / {a.workers} workers = {proj / 3600:.2f} h")  # fmt: skip
            if proj > BUDGET_S:
                keep_n = i + int(BUDGET_S * a.workers / (2 * med))
                print(f"projection over {BUDGET_S / 3600:.0f} h: drafting {keep_n} of {len(todo)} targets")
                todo = todo[:keep_n]
        results += list(ex.map(work, todo[i:]))
    rows = sorted((r["row"] for r in results if "row" in r), key=lambda r: int(r["qid"].split(":")[1]))
    rejects = [{k: v for k, v in r.items() if k != "row"} for r in results if "row" not in r]
    write_jsonl(OUT if full else CACHE / "pilot.jsonl", rows)
    write_jsonl(CACHE / ("rejects.jsonl" if full else "pilot_rejects.jsonl"), rejects)
    report(rows, rejects, grok, time.monotonic() - t_start, OUT if full else CACHE / "pilot.jsonl")
    return 0


def report(rows: list[dict], rejects: list[dict], grok: Grok, seconds: float, out: Path) -> None:
    kept = [r for r in rows if r["meta"]["status"] == "keep"]
    print(f"\nwrote {out}: {len(rows)} rows, {len(kept)} kept, {len(rows) - len(kept)} dropped in review, "
          f"{len(rejects)} failed the automatic checks")  # fmt: skip
    for name, key in (("format", lambda r: r["format"]), ("type", lambda r: r["meta"]["type"]),
                      ("split", lambda r: r["meta"]["split"]), ("lookalike", lambda r: r["meta"]["lookalike"])):  # fmt: skip
        print(f"kept by {name}: {dict(sorted(Counter(map(key, kept)).items(), key=str))}")
    print(f"automatic-check rejects: {dict(Counter(r['reject'] for r in rejects).most_common())}")
    print(f"review drops by format: {dict(Counter(r['format'] for r in rows if r['meta']['status'] == 'drop'))}")
    if grok.latency:
        print(f"grok live calls {len(grok.latency)}: median {statistics.median(grok.latency):.1f}s, "
              f"p95 {pct(grok.latency, 0.95):.1f}s, sum {sum(grok.latency) / 60:.1f} min")  # fmt: skip
    print(f"wall time {seconds / 60:.1f} min")
    held = sorted(r["qid"] for r in kept if r["meta"]["split"] == "heldout")
    print(
        f"heldout kept {len(held)}; sha256 of sorted qids {hashlib.sha256(chr(10).join(held).encode()).hexdigest()[:16]}"
    )


if __name__ == "__main__":
    sys.exit(main())
