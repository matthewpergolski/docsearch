"""File metadata from inside each file (title, author, dates) plus its name and folder, and the search filter.

One JSON row per doc_id in `meta.jsonl` next to the index folder (work/meta.jsonl by default). Dates come from the
file's own properties (Office core properties, the PDF Info dictionary), never from the file system: a copied
folder gives every file the copy time, and a wrong date is worse than none. An empty date means unknown.
"""

from __future__ import annotations

import json
import re
from collections.abc import Callable
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime
from pathlib import Path, PurePosixPath

from . import ui
from .config import Config
from .manifest import Doc, read_manifest

Props = tuple[str | None, str | None, str, str]  # title, author, created, modified ("YYYY-MM-DD" or "")


@dataclass(frozen=True)
class DocMeta:
    doc_id: str
    path: str  # relative to corpus_dir; a file synced from outside corpus_dir: its doc_id ("<folder name>/sub/file")
    format: str
    title: str = ""
    author: str = ""
    created: str = ""
    modified: str = ""
    date_source: str = "none"  # "file properties" | "pdf metadata" | "none"
    folder: str = ""  # the folder of `path`; "" at the top
    file_name: str = ""

    @property
    def date(self) -> str:
        """The date a filter and the hit line use: the last change, else the creation."""
        return self.modified or self.created


def _day(dt: datetime | None) -> str:
    return dt.date().isoformat() if dt else ""


def _docx(p: Path) -> Props:
    from docx import Document

    c = Document(str(p)).core_properties
    return c.title, c.author, _day(c.created), _day(c.modified)


def _pptx(p: Path) -> Props:
    from pptx import Presentation

    c = Presentation(str(p)).core_properties
    return c.title, c.author, _day(c.created), _day(c.modified)


def _xlsx(p: Path) -> Props:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True)
    try:
        c = wb.properties
    finally:
        wb.close()
    return c.title, c.creator, _day(c.created), _day(c.modified)


_PDF_DATE = re.compile(r"^(?:D:)?(\d{4})(\d{2})?(\d{2})?")


def pdf_date(s: str) -> str:
    """A PDF date string (D:YYYYMMDDHHmmSS+HH'mm', every part after the year optional) as YYYY-MM-DD in the time
    of day the file states; a missing month or day is 01, as the PDF standard says. "" when it does not parse."""
    hit = _PDF_DATE.match(s.strip())
    if not hit:
        return ""
    y, mo, d = hit.group(1), hit.group(2) or "01", hit.group(3) or "01"
    try:
        return date(int(y), int(mo), int(d)).isoformat()
    except ValueError:
        return ""


def _pdf(p: Path) -> Props:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(p))
    try:
        m = doc.get_metadata_dict()
    finally:
        doc.close()
    return m.get("Title"), m.get("Author"), pdf_date(m.get("CreationDate", "")), pdf_date(m.get("ModDate", ""))


# format -> (date_source, reader). Formats not listed (csv, image, slide_image) carry no properties: names only.
EXTRACTORS: dict[str, tuple[str, Callable[[Path], Props]]] = {
    "docx": ("file properties", _docx),
    "pptx": ("file properties", _pptx),
    "xlsx": ("file properties", _xlsx),
    "pdf": ("pdf metadata", _pdf),
}
# Read as their format but their properties are not: a .doc holds them in an OLE stream python-docx cannot open,
# and both converters change its dates (soffice moves them by the local UTC offset, textutil drops the modified
# date; DOC1), so a .doc carries its name and folder only.
NO_PROPERTIES = {".doc"}


def extract(corpus: Path, d: Doc) -> DocMeta:
    p = PurePosixPath(d.path)
    if p.is_absolute():  # synced from a folder outside corpus_dir: never show the absolute path
        p = PurePosixPath(d.doc_id)
    folder = "" if str(p.parent) == "." else p.parent.as_posix()
    base = DocMeta(d.doc_id, p.as_posix(), d.format, folder=folder, file_name=p.name)
    if d.format not in EXTRACTORS or p.suffix.lower() in NO_PROPERTIES:
        return base
    source, read = EXTRACTORS[d.format]
    try:
        title, author, created, modified = read(corpus / d.path)
    except Exception as e:  # noqa: BLE001 - a damaged or locked file still gets its name and folder
        ui.logger.warning("meta: cannot read properties of %s: %s: %s", d.doc_id, type(e).__name__, e)
        return base
    return replace(base, title=(title or "").strip(), author=(author or "").strip(), created=created,
                   modified=modified, date_source=source if created or modified else "none")  # fmt: skip


def meta_path(cfg: Config) -> Path:
    return cfg.path("index_dir").parent / "meta.jsonl"


def read_meta(cfg: Config) -> dict[str, DocMeta]:
    p = meta_path(cfg)
    if not p.exists():
        return {}
    return {r["doc_id"]: DocMeta(**r) for r in map(json.loads, p.read_text().splitlines()) if r}


def _write(cfg: Config, rows: dict[str, DocMeta]) -> None:
    p = meta_path(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".jsonl.tmp")
    tmp.write_text("".join(json.dumps(asdict(rows[k]), ensure_ascii=False) + "\n" for k in sorted(rows)))
    tmp.replace(p)


def update_meta(cfg: Config, rows: list[DocMeta], drop: set[str] = frozenset()) -> None:
    """Rewrite these docs' rows and remove the dropped ones; every other row stays."""
    cur = {k: v for k, v in read_meta(cfg).items() if k not in drop}
    _write(cfg, cur | {r.doc_id: r for r in rows})


def build_meta(cfg: Config) -> dict:
    """One row for every doc in corpus/manifest.csv; rows of docs no longer listed go."""
    corpus = cfg.path("corpus_dir")
    docs = read_manifest(corpus / "manifest.csv")
    rows = {d.doc_id: extract(corpus, d) for d in docs}
    _write(cfg, rows)
    return {"files": len(rows), "with_title": sum(bool(r.title) for r in rows.values()),
            "with_date": sum(bool(r.date) for r in rows.values()), "path": str(meta_path(cfg))}  # fmt: skip


@dataclass(frozen=True)
class Filter:
    """What the user narrows a search to. Dates are YYYY-MM-DD and both ends are inclusive. When after or before
    is set, a file with no date of its own is left out (its date is unknown, not inside the range) unless undated
    is set; note() says how many such files a search left out."""

    formats: tuple[str, ...] = ()
    after: str | None = None
    before: str | None = None
    folder: str | None = None  # a folder relative to corpus_dir (or a synced folder's name); subfolders count too
    undated: bool = False  # keep files with no date of their own when after or before is set

    @property
    def active(self) -> bool:
        return bool(self.formats) or any(x is not None for x in (self.after, self.before, self.folder))

    def allows(self, m: DocMeta) -> bool:
        if self.formats and m.format not in self.formats:
            return False
        if self.folder is not None:
            top = self.folder.strip("/")
            if top and m.folder != top and not m.folder.startswith(top + "/"):
                return False
        if self.after is not None or self.before is not None:
            if not m.date:
                return self.undated
            if self.after is not None and m.date < self.after:
                return False
            if self.before is not None and m.date > self.before:
                return False
        return True

    def note(self, metas) -> str | None:
        """How many files a date filter left out only because they have no date; None when that is none."""
        if self.undated or (self.after is None and self.before is None):
            return None
        n = sum(1 for m in metas if not m.date and replace(self, after=None, before=None).allows(m))
        if not n:
            return None
        one = n == 1
        return (
            f"date filter: {n} file{'' if one else 's'} {'has' if one else 'have'} no date of {'its' if one else 'their'} "
            f"own and {'was' if one else 'were'} left out (add --undated to keep {'it' if one else 'them'})"
        )
