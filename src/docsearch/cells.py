"""Phase C: table cells as one long table (work/cells.sqlite, built by scripts/build_cells.py).

A grid is a table as rows of cell texts, spans repeated into every position they cover and "" for an empty
position. Rows and columns are 1-based. A cell carries a number only when the whole cell is one number; a cell
with two or more numbers and no word is fused (a parser merged neighbouring cells)."""

from __future__ import annotations

import re
import sqlite3
from datetime import date, datetime, time
from functools import lru_cache
from pathlib import Path

SCHEMA = """
CREATE TABLE tables (table_id TEXT PRIMARY KEY, doc_id TEXT, format TEXT, page INTEGER, sheet TEXT,
                     n_rows INTEGER, n_cols INTEGER, quality REAL, caption TEXT);
CREATE TABLE cells (table_id TEXT, doc_id TEXT, format TEXT, page INTEGER, sheet TEXT, row INTEGER, col INTEGER,
                    a1 TEXT, text TEXT, num REAL, row_label TEXT, col_header TEXT, section TEXT, fused INTEGER);
CREATE INDEX cells_doc_page ON cells (doc_id, page);
"""

MISSING = {":", "-", "–", "—", ".."}
_NUM = re.compile(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d{1,3}(?:[   ]\d{3})+|\d+)(?:\.\d+)?|[+-]?\.\d+")
_TOKEN = re.compile(r"\d+(?:[.,]\d+)*")
_WORD = re.compile(r"[^\W\d_]+")
UNITS = {"pps", "pp", "pts", "pt", "bn", "mn", "m", "k", "p", "e", "b", "f", "u", "r", "c", "s", "d", "z", "i"}
_SECTION = re.compile(r"^\d{1,3}\s*[-–—.)]\s*[^\W\d_(]")
_YEAR = re.compile(r"(19|20)\d\d")


def parse_num(text: str) -> float | None:
    """'(569)' -> -569, '1,320' -> 1320, '4.5%' -> 4.5, '£44,431' -> 44431; anything else -> None."""
    s = (text or "").strip().replace("−", "-")
    neg = s.startswith("(") and s.endswith(")")
    s = s[1:-1].strip() if neg else s
    s = s.lstrip("£$€¥").strip().removesuffix("%").strip()
    if not _NUM.fullmatch(s):
        return None
    v = float(re.sub(r"[,   ]", "", s))
    return -v if neg else v


def is_fused(text: str) -> bool:
    """Two or more numbers and no word other than units or one-letter flags ('7.5 p 6.1 e', '0.1 pps 3.4 %')."""
    return len(_TOKEN.findall(text or "")) >= 2 and all(w.lower() in UNITS for w in _WORD.findall(text))


def cell_numbers(text: str, num: float | None, fused: bool) -> set[float]:
    """The numbers a cell holds: its own number, or each number of a fused cell."""
    if num is not None:
        return {num}
    if fused:
        return {float(t.replace(",", "")) for t in _TOKEN.findall(text) if re.fullmatch(r"\d+(?:,\d{3})*(?:\.\d+)?", t)}
    return set()


def _kind(text: str) -> str:
    if not text:
        return "empty"
    if text in MISSING:
        return "missing"
    return "num" if parse_num(text) is not None else "text"


def _year(text: str) -> bool:
    return bool(_YEAR.fullmatch(text.strip()))


def grid_cells(grid: list[list[str]]) -> list[dict]:
    """The non-empty cells of a grid with row label, column header and section. Header rows are the leading rows
    whose only numbers are years; a table with no other number takes its first non-empty row as header."""
    rows = [[(t or "").strip() for t in r] for r in grid]
    kinds = [[_kind(t) for t in r] for r in rows]
    first_data = next((i for i, k in enumerate(kinds)
                       if any(x == "num" and not _year(t) for x, t in zip(k, rows[i], strict=True))), None)  # fmt: skip
    if first_data is None:
        first_data = next((i + 1 for i, r in enumerate(rows) if any(r)), len(rows))
    width = max((len(r) for r in rows), default=0)
    headers = []
    for c in range(width):
        parts: list[str] = []
        for r in rows[:first_data]:
            t = r[c] if c < len(r) else ""
            if t and (not parts or parts[-1] != t):
                parts.append(t)
        headers.append(" ".join(parts))
    out, section = [], None
    for i, (r, k) in enumerate(zip(rows, kinds, strict=True)):
        label = next((t for t, x in zip(r, k, strict=True) if x == "text" and re.search(r"[^\W\d_]", t)), None)
        data = i >= first_data
        if data and label and _SECTION.match(label):
            section = label
        for j, (t, x) in enumerate(zip(r, k, strict=True)):
            if x == "empty":
                continue
            num = parse_num(t) if x == "num" else None
            out.append({"row": i + 1, "col": j + 1, "text": t, "num": num, "fused": x == "text" and is_fused(t),
                        "row_label": label, "col_header": (headers[j] or None) if data else None,
                        "section": section if data else None})  # fmt: skip
    return out


def table_quality(grid: list[list[str]]) -> tuple[int, int, float]:
    """(rows, columns, share of grid positions that hold text)."""
    n_rows, n_cols = len(grid), max((len(r) for r in grid), default=0)
    filled = sum(1 for r in grid for t in r if (t or "").strip())
    return n_rows, n_cols, (filled / (n_rows * n_cols) if n_rows and n_cols else 0.0)


def docling_grid(data: dict) -> list[list[str]]:
    """A Docling table's cells with row and column spans repeated into every position they cover."""
    grid = [[""] * int(data["num_cols"]) for _ in range(int(data["num_rows"]))]
    for c in data.get("table_cells") or []:
        for r in range(c["start_row_offset_idx"], c["end_row_offset_idx"]):
            for k in range(c["start_col_offset_idx"], c["end_col_offset_idx"]):
                if r < len(grid) and k < len(grid[r]):
                    grid[r][k] = c.get("text") or ""
    return grid


def xlsx_text(value, number_format: str) -> str:
    """A cell value as the sheet shows it: percentages times 100, dates as ISO, floats without binary noise."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, (int, float)):
        if "%" in (number_format or ""):
            return f"{value * 100:.10g}%"
        return f"{value:.10g}" if isinstance(value, float) else str(value)
    return str(value).strip()


def xlsx_sheets(src) -> tuple[list[tuple[int, str, list[list[str]]]], list[str]]:
    """(page, sheet name, grid) per visible worksheet, merged ranges repeated, and the names of the hidden sheets
    that were skipped. Page = position in the workbook, hidden sheets counted, as Docling numbers sheets."""
    from openpyxl import load_workbook

    wb = load_workbook(src, data_only=True)
    out, skipped = [], []
    for page, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        if ws.sheet_state != "visible":
            skipped.append(name)
            continue
        if hasattr(ws, "iter_rows"):
            out.append((page, name, _sheet_grid(ws)))
    return out, skipped


def _sheet_grid(ws) -> list[list[str]]:
    """A worksheet from A1, each merged range's value repeated over the range."""
    merged = {}
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = top

    def text(c) -> str:
        c = merged.get((c.row, c.column), c)
        return xlsx_text(c.value, c.number_format)

    return _trim([[text(c) for c in row] for row in ws.iter_rows()])


def csv_grid(text: str) -> list[list[str]]:
    import csv

    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _trim([[t.strip() for t in r] for r in csv.reader(text.splitlines(), dialect)])


def _trim(grid: list[list[str]]) -> list[list[str]]:
    """Cut empty rows and columns at the bottom and right; pad rows to one width."""
    rows = list(grid)
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((j + 1 for j, t in enumerate(r) if t), default=0) for r in rows), default=0)
    return [(r + [""] * width)[:width] for r in rows]


# -- C2 table lookup: which page of a file holds the row and column a query names

STOP = {"the", "and", "for", "was", "were", "are", "did", "does", "how", "what", "which", "when", "where", "who",
        "whose", "with", "from", "that", "this", "than", "into", "many", "much", "their", "there", "they", "have",
        "has", "had", "been", "its", "per", "give", "given", "show", "shown", "report", "reported", "about", "during",
        "between", "rate", "year", "years", "percent", "percentage", "points", "change", "value", "figure", "total"}  # fmt: skip
_TERM = re.compile(r"\d+(?:\.\d+)?|[^\W\d_]{3,}")
_QVAL = re.compile(r"\d[\d,]*(?:\.\d+)?")


def terms(text: str | None) -> frozenset[str]:
    """Lower-case words of 3+ letters (stop words dropped) and numbers; '15-24' gives '15' and '24'."""
    return frozenset(t for t in _TERM.findall((text or "").lower()) if t not in STOP)


def query_values(query: str) -> set[float]:
    """Numbers a query may hold as a cell value: decimals, and integers of 100 or more that are not years."""
    out = set()
    for t in _QVAL.findall(query):
        t = t.rstrip(",")
        v = float(t.replace(",", ""))
        year = "," not in t and _YEAR.fullmatch(t)
        if "." in t or (v >= 100 and not year):
            out.add(v)
    return out


def page_rows(cells: list[dict]) -> dict[int, list[tuple]]:
    """Per page, one (label terms, section terms, [(header terms, numbers)]) per table row."""
    rows: dict[tuple, list[dict]] = {}
    for c in cells:
        rows.setdefault((c["page"], c["table_id"], c["row"]), []).append(c)
    out: dict[int, list[tuple]] = {}
    for (page, _, _), cs in rows.items():
        out.setdefault(page, []).append((terms(cs[0]["row_label"]), terms(cs[0]["section"]),
                                         [(terms(c["col_header"]), cell_numbers(c["text"], c["num"], bool(c["fused"])))
                                          for c in cs]))  # fmt: skip
    return out


def score_rows(query: str, rows_by_page: dict[int, list[tuple]]) -> dict[int, int]:
    """A row matches when a query value sits in it, or when its label and a column header share query terms.
    Strength: 2 per query value found + label terms + header terms (of the cells holding a value, else the best
    cell) + section terms. A page scores its strongest row; pages with no match are left out."""
    q, vals = terms(query), query_values(query)
    out: dict[int, int] = {}
    for page, rows in rows_by_page.items():
        for label_t, section_t, cells in rows:
            label = len(label_t & q)
            hit = [(h, n & vals) for h, n in cells if n & vals]
            found = set().union(*(n for _, n in hit)) if hit else set()
            head = max((len(h & q) for h, _ in (hit or cells)), default=0)
            if found or (label and head):
                out[page] = max(out.get(page, 0), 2 * len(found) + label + head + len(section_t & q))
    return out


def page_scores(query: str, cells: list[dict]) -> dict[int, int]:
    return score_rows(query, page_rows(cells))


class CellLookup:
    """work/cells.sqlite opened read-only; the rows of each (doc, page) are read once and kept."""

    COLS = "page, table_id, row, text, num, fused, row_label, col_header, section"

    def __init__(self, path: Path):
        if not Path(path).exists():
            raise SystemExit(f"--table-lookup needs {path}; build it with: uv run python scripts/build_cells.py")
        self.db = sqlite3.connect(f"file:{path}?mode=ro", uri=True, check_same_thread=False)
        self._rows = lru_cache(maxsize=1024)(self._read)

    def _read(self, doc: str, page: int) -> list[tuple]:
        cur = self.db.execute(f"SELECT {self.COLS} FROM cells WHERE doc_id = ? AND page = ?", (doc, page))
        names = [d[0] for d in cur.description]
        return page_rows([dict(zip(names, r, strict=True)) for r in cur]).get(page, [])

    def scores(self, query: str, doc: str, pages: list[int]) -> dict[int, int]:
        return score_rows(query, {p: self._rows(doc, p) for p in pages})


def cells_path(cfg) -> Path:
    return cfg.path("parsed_dir").parent / "cells.sqlite"
