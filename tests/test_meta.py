"""Phase A: file metadata from inside each file, the search filter, and the allow-list in retrieval."""

import io
import json
import re
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import numpy as np

from docsearch.config import load_config
from docsearch.manifest import Doc, write_manifest
from docsearch.meta import DocMeta, Filter, build_meta, extract, meta_path, pdf_date, read_meta, update_meta


def tiny_pdf(path: Path, info: dict[str, str]) -> None:
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>",
            b"<< " + b" ".join(f"/{k} ({v})".encode() for k, v in info.items()) + b" >>"]  # fmt: skip
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R /Info 4 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def set_xlsx_modified(path: Path, iso: str) -> None:
    """openpyxl stamps the save time as the modified date; write a fixed one into docProps/core.xml."""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + iso.encode(), data)
            z.writestr(info, data)


class Tmp(unittest.TestCase):
    def setUp(self):
        self.td = tempfile.TemporaryDirectory()
        self.root = Path(self.td.name)
        (self.root / "docsearch.yaml").write_text("subset: {enabled: false}\nvisual: {default: null}\n")
        self.cfg = load_config(self.root / "docsearch.yaml", env={})
        self.corpus = self.cfg.path("corpus_dir")
        (self.corpus / "sub").mkdir(parents=True)

    def tearDown(self):
        self.td.cleanup()


class TestExtract(Tmp):
    def test_docx_core_properties(self):
        from docx import Document

        doc = Document()
        doc.core_properties.title = "Q3 plan"
        doc.core_properties.author = "Ada"
        doc.core_properties.created = datetime(2024, 1, 31, 10, 0)
        doc.core_properties.modified = datetime(2024, 2, 15, 8, 30)
        doc.save(self.corpus / "sub" / "plan.docx")
        got = extract(self.corpus, Doc("s/sub/plan.docx", "sub/plan.docx", "docx", "s", ""))
        self.assertEqual(got, DocMeta("s/sub/plan.docx", "sub/plan.docx", "docx", "Q3 plan", "Ada", "2024-01-31",
                                      "2024-02-15", "file properties", "sub", "plan.docx"))  # fmt: skip
        self.assertEqual(got.date, "2024-02-15")

    def test_pptx_core_properties(self):
        from pptx import Presentation

        prs = Presentation()
        prs.core_properties.title = "Board deck"
        prs.core_properties.author = "Lin"
        prs.core_properties.created = datetime(2023, 6, 1, 12, 0)
        prs.core_properties.modified = datetime(2023, 6, 2, 12, 0)
        prs.save(self.corpus / "deck.pptx")
        got = extract(self.corpus, Doc("s/deck.pptx", "deck.pptx", "pptx", "s", ""))
        self.assertEqual(got, DocMeta("s/deck.pptx", "deck.pptx", "pptx", "Board deck", "Lin", "2023-06-01",
                                      "2023-06-02", "file properties", "", "deck.pptx"))  # fmt: skip

    def test_xlsx_workbook_properties(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.properties.title = "Budget 2026"
        wb.properties.creator = "Sam"
        wb.properties.created = datetime(2026, 1, 5, 9, 0)
        wb.save(self.corpus / "budget.xlsx")
        set_xlsx_modified(self.corpus / "budget.xlsx", "2026-01-06T09:00:00Z")
        got = extract(self.corpus, Doc("s/budget.xlsx", "budget.xlsx", "xlsx", "s", ""))
        self.assertEqual(got, DocMeta("s/budget.xlsx", "budget.xlsx", "xlsx", "Budget 2026", "Sam", "2026-01-05",
                                      "2026-01-06", "file properties", "", "budget.xlsx"))  # fmt: skip

    def test_pdf_metadata(self):
        tiny_pdf(self.corpus / "sub" / "r.pdf", {"Title": "Annual report", "Author": "Kim",
                                                  "CreationDate": "D:20240131120000+01'00'", "ModDate": "D:2024"})  # fmt: skip
        got = extract(self.corpus, Doc("s/sub/r.pdf", "sub/r.pdf", "pdf", "s", ""))
        self.assertEqual(got, DocMeta("s/sub/r.pdf", "sub/r.pdf", "pdf", "Annual report", "Kim", "2024-01-31",
                                      "2024-01-01", "pdf metadata", "sub", "r.pdf"))  # fmt: skip

    def test_pdf_without_dates_has_no_date(self):
        tiny_pdf(self.corpus / "r.pdf", {"Title": "Untitled scan"})
        got = extract(self.corpus, Doc("s/r.pdf", "r.pdf", "pdf", "s", ""))
        self.assertEqual((got.title, got.date, got.date_source), ("Untitled scan", "", "none"))

    def test_csv_and_image_get_names_only(self):
        (self.corpus / "sub" / "list.csv").write_text("a,b\n1,2\n")
        got = extract(self.corpus, Doc("s/sub/list.csv", "sub/list.csv", "csv", "s", ""))
        self.assertEqual(got, DocMeta("s/sub/list.csv", "sub/list.csv", "csv", "", "", "", "", "none", "sub",
                                      "list.csv"))  # fmt: skip

    def test_damaged_file_keeps_its_name_and_folder(self):
        (self.corpus / "sub" / "bad.docx").write_bytes(b"not a zip")
        got = extract(self.corpus, Doc("s/sub/bad.docx", "sub/bad.docx", "docx", "s", ""))
        self.assertEqual(got, DocMeta("s/sub/bad.docx", "sub/bad.docx", "docx", "", "", "", "", "none", "sub",
                                      "bad.docx"))  # fmt: skip

    def test_old_word_file_gets_names_only_and_no_warning(self):
        (self.corpus / "sub" / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0 old word")
        (self.corpus / "sub" / "bad.docx").write_bytes(b"\xd0\xcf\x11\xe0 old word")
        with self.assertNoLogs("docsearch", "WARNING"):
            got = extract(self.corpus, Doc("s/sub/old.doc", "sub/old.doc", "docx", "s", ""))
        self.assertEqual(got, DocMeta("s/sub/old.doc", "sub/old.doc", "docx", "", "", "", "", "none", "sub",
                                      "old.doc"))  # fmt: skip
        with self.assertLogs("docsearch", "WARNING"):  # the same bytes named .docx are a damaged docx
            extract(self.corpus, Doc("s/sub/bad.docx", "sub/bad.docx", "docx", "s", ""))

    def test_file_synced_from_outside_corpus_dir_is_named_by_its_doc_id(self):
        outside = self.root / "home" / "files" / "sub"
        outside.mkdir(parents=True)
        (outside / "list.csv").write_text("a,b\n")
        got = extract(self.corpus, Doc("files/sub/list.csv", str(outside / "list.csv"), "csv", "files", ""))
        self.assertEqual(got, DocMeta("files/sub/list.csv", "files/sub/list.csv", "csv", folder="files/sub",
                                      file_name="list.csv"))  # fmt: skip
        self.assertTrue(Filter(folder="files").allows(got))
        self.assertNotIn(str(self.root), json.dumps(got.__dict__))


class TestPdfDate(unittest.TestCase):
    def test_pdf_date_strings(self):
        cases = {
            "D:20240131120000+01'00'": "2024-01-31",
            "D:20240131120000Z": "2024-01-31",
            "D:202402": "2024-02-01",
            "D:2024": "2024-01-01",
            "20240131": "2024-01-31",
            "D:20241341": "",
            "Wed Jan 31 2024": "",
            "": "",
        }
        self.assertEqual({s: pdf_date(s) for s in cases}, cases)


def m(fmt="pdf", date="", folder="", created=""):
    return DocMeta("s/x", "x", fmt, modified=date, created=created, folder=folder,
                   date_source="pdf metadata" if date or created else "none")  # fmt: skip


class TestFilter(unittest.TestCase):
    def test_allows_truth_table(self):
        cases = [
            (Filter(), m(), True),
            (Filter(formats=("xlsx",)), m("xlsx"), True),
            (Filter(formats=("xlsx",)), m("pdf"), False),
            (Filter(formats=("xlsx", "pdf")), m("pdf"), True),
            (Filter(after="2024-01-01"), m(date="2024-01-01"), True),
            (Filter(after="2024-01-01"), m(date="2023-12-31"), False),
            (Filter(after="2024-01-01"), m(), False),
            (Filter(after="2024-01-01"), m(created="2024-05-01"), True),
            (Filter(before="2024-01-01"), m(date="2024-01-01"), True),
            (Filter(before="2024-01-01"), m(date="2024-01-02"), False),
            (Filter(before="2024-01-01"), m(), False),
            (Filter(after="2024-01-01", before="2024-12-31"), m(date="2024-06-30"), True),
            (Filter(after="2024-01-01", before="2024-12-31"), m(date="2025-01-01"), False),
            (Filter(folder="reports"), m(folder="reports"), True),
            (Filter(folder="reports"), m(folder="reports/2024"), True),
            (Filter(folder="reports/"), m(folder="reports/2024"), True),
            (Filter(folder="reports"), m(folder="reports-old"), False),
            (Filter(folder="reports"), m(folder=""), False),
            (Filter(formats=("pdf",), folder="reports"), m("xlsx", folder="reports"), False),
        ]
        got = [(f, x, f.allows(x)) for f, x, _ in cases]
        self.assertEqual(got, cases)

    def test_undated_keeps_files_with_no_date(self):
        cases = [
            (Filter(after="2024-01-01", undated=True), m(), True),
            (Filter(before="2024-01-01", undated=True), m(), True),
            (Filter(after="2024-01-01", undated=True), m(date="2023-12-31"), False),
            (Filter(formats=("pdf",), after="2024-01-01", undated=True), m("csv"), False),
        ]
        self.assertEqual([(f, x, f.allows(x)) for f, x, _ in cases], cases)

    def test_note_counts_undated_files_a_date_filter_left_out(self):
        metas = [m(), m("csv"), m(date="2024-06-01"), m(date="2020-01-01"), m("xlsx", folder="other")]
        self.assertEqual(Filter(after="2024-01-01").note(metas),
                         "date filter: 3 files have no date of their own and were left out "
                         "(add --undated to keep them)")  # fmt: skip
        self.assertEqual(Filter(after="2024-01-01", folder="other").note(metas),
                         "date filter: 1 file has no date of its own and was left out "
                         "(add --undated to keep it)")  # fmt: skip
        self.assertIsNone(Filter(after="2024-01-01", undated=True).note(metas))
        self.assertIsNone(Filter(formats=("csv",)).note(metas))
        self.assertIsNone(Filter(after="2024-01-01").note([m(date="2024-06-01")]))

    def test_active(self):
        self.assertEqual([Filter().active, Filter(formats=("csv",)).active, Filter(after="2024-01-01").active,
                          Filter(folder="a").active], [False, True, True, True])  # fmt: skip


class TestStore(Tmp):
    def test_build_update_and_drop_converge(self):
        (self.corpus / "a.csv").write_text("x\n")
        (self.corpus / "sub" / "b.csv").write_text("y\n")
        write_manifest(self.corpus / "manifest.csv", [Doc("s/a.csv", "a.csv", "csv", "s", ""),
                                                       Doc("s/sub/b.csv", "sub/b.csv", "csv", "s", "")])  # fmt: skip
        self.assertEqual(build_meta(self.cfg)["files"], 2)
        self.assertEqual(meta_path(self.cfg), self.root.resolve() / "work" / "meta.jsonl")
        first = meta_path(self.cfg).read_text()
        build_meta(self.cfg)
        self.assertEqual(meta_path(self.cfg).read_text(), first)
        self.assertEqual(sorted(read_meta(self.cfg)), ["s/a.csv", "s/sub/b.csv"])
        self.assertEqual(read_meta(self.cfg)["s/sub/b.csv"].folder, "sub")
        row = DocMeta("s/a.csv", "a.csv", "csv", "Renamed", file_name="a.csv")
        update_meta(self.cfg, [row])
        update_meta(self.cfg, [row])
        self.assertEqual(read_meta(self.cfg)["s/a.csv"].title, "Renamed")
        self.assertEqual(len(meta_path(self.cfg).read_text().splitlines()), 2)
        update_meta(self.cfg, [], drop={"s/sub/b.csv"})
        self.assertEqual(sorted(read_meta(self.cfg)), ["s/a.csv"])

    def test_sync_ops_write_and_drop_rows(self):
        from docsearch.sync import Ops

        (self.corpus / "a.csv").write_text("x\n")
        ops = Ops(self.cfg)
        ops.meta(Doc("s/a.csv", "a.csv", "csv", "s", ""))
        ops.flush_meta()
        self.assertEqual(read_meta(self.cfg)["s/a.csv"].file_name, "a.csv")
        ops.remove_artifacts({"s/a.csv"})
        self.assertEqual(read_meta(self.cfg), {})

    def test_sync_ops_write_meta_once_per_sync(self):
        import docsearch.meta as meta_mod
        from docsearch.sync import Ops

        docs = [Doc(f"s/{n}.csv", f"{n}.csv", "csv", "s", "") for n in "abc"]
        for d in docs:
            (self.corpus / d.path).write_text("x\n")
        ops = Ops(self.cfg)
        with mock.patch.object(meta_mod, "_write", wraps=meta_mod._write) as w:
            for d in docs:
                ops.meta(d)
            self.assertEqual(read_meta(self.cfg), {})
            ops.flush_meta()
            ops.flush_meta()
        self.assertEqual(w.call_count, 1)
        self.assertEqual(sorted(read_meta(self.cfg)), ["s/a.csv", "s/b.csv", "s/c.csv"])


N = 300


def lance_table(cfg):
    """300 chunks, one per doc d0..d299, all of format pdf except d150 (xlsx). Both BM25 and the vector rank the
    docs in id order, so d150 ranks 151st without a filter."""
    from docsearch.index import connect

    rows = [{"chunk_id": f"c{i}", "doc_id": f"d{i}", "format": "xlsx" if i == 150 else "pdf", "source": "s",
             "pages": [1], "text": "apple " * (N - i) + f"tail{i}",
             "vector": np.array([1.0, i / N, 0.0, 0.0], dtype=np.float32)} for i in range(N)]  # fmt: skip
    tbl = connect(cfg).create_table("chunks_hybrid", rows)
    tbl.create_fts_index("text", use_tantivy=False)
    return tbl


QVEC = np.array([1.0, 0.0, 0.0, 0.0], dtype=np.float32)


class TestAllowList(Tmp):
    def setUp(self):
        super().setUp()
        lance_table(self.cfg)
        from docsearch.search import Retriever

        self.r = Retriever(self.cfg, "hybrid")

    def test_no_filter_path_returns_the_unfiltered_top_k(self):
        self.assertEqual((self.r.sc.bm25_k, self.r.sc.vector_k), (100, 100))
        ids = [c["chunk_id"] for c in self.r.candidates("apple", QVEC)]
        self.assertEqual(ids, [f"c{i}" for i in range(100)])

    def test_allowed_doc_ranked_151st_still_comes_back(self):
        self.assertNotIn("c150", [r["chunk_id"] for r in self.r.bm25("apple", 100)])
        self.assertNotIn("c150", [r["chunk_id"] for r in self.r.vector(QVEC, 100)])
        allow = frozenset({"d150"})
        with self.assertNoLogs("docsearch", "WARNING"):
            self.assertEqual([r["chunk_id"] for r in self.r.bm25("apple", 100, allow)], ["c150"])
        self.assertEqual([r["chunk_id"] for r in self.r.vector(QVEC, 100, allow)], ["c150"])
        self.assertEqual([c["chunk_id"] for c in self.r.candidates("apple", QVEC, allow=allow)], ["c150"])

    def test_allow_list_keeps_k(self):
        allow = frozenset(f"d{i}" for i in range(150, 300))
        ids = [c["chunk_id"] for c in self.r.candidates("apple", QVEC, allow=allow)]
        self.assertEqual(ids, [f"c{i}" for i in range(150, 250)])

    def test_doc_id_with_quote_is_escaped(self):
        with self.assertNoLogs("docsearch", "WARNING"):
            self.assertEqual(self.r.bm25("apple", 10, frozenset({"it's"})), [])

    def test_fts_with_prefilter_and_in_list_runs_on_the_installed_lancedb(self):
        with self.assertNoLogs("docsearch", "WARNING"):
            got = self.r.bm25("apple", 100, frozenset({"d3", "d150", "d299"}))
        self.assertEqual([r["chunk_id"] for r in got], ["c3", "c150", "c299"])

    def test_a_failing_bm25_query_logs_a_warning_and_gives_no_hits(self):
        with mock.patch("docsearch.search._doc_in", side_effect=ValueError("bad filter")), \
                self.assertLogs("docsearch", "WARNING") as logs:  # fmt: skip
            self.assertEqual(self.r.bm25("apple", 10, frozenset({"d1"})), [])
        self.assertIn("bm25", logs.output[0])
        self.assertIn("bad filter", logs.output[0])
        with mock.patch.object(self.r.tbl, "search", side_effect=RuntimeError("fts index missing")), \
                self.assertLogs("docsearch", "WARNING") as logs:  # fmt: skip
            self.assertEqual(self.r.bm25("apple", 10), [])
        self.assertIn("fts index missing", logs.output[0])


class FakeStore:
    def __init__(self, units):
        self.ids = units

    def search(self, q, k):
        return [(u, 1.0 - i / 1000) for i, u in enumerate(self.ids[:k])]


def fake_ranker():
    from docsearch.visual import VisualRanker

    units = [f"d{i}#p1" for i in range(N)] + ["scan#p1"]
    vr = object.__new__(VisualRanker)
    vr.store = FakeStore(units)
    vr.meta = {u: {"unit": u, "doc_id": u.split("#")[0], "format": "pdf", "page": 1} for u in units}
    vr.docs = {x["doc_id"] for x in vr.meta.values()}
    return vr


class TestVisualAllowList(Tmp):
    def test_rank_skips_disallowed_units_before_cutting_to_k(self):
        vr = fake_ranker()
        self.assertEqual([u for u, _ in vr.rank(QVEC, 3)], ["d0#p1", "d1#p1", "d2#p1"])
        self.assertEqual([u for u, _ in vr.rank(QVEC, 3, allow=frozenset({"d150", "scan"}))],
                         ["d150#p1", "scan#p1"])  # fmt: skip

    def test_format_aware_fusion_with_allow_list(self):
        from docsearch.search import Retriever

        lance_table(self.cfg)
        r = Retriever(self.cfg, "hybrid")
        r.visual = fake_ranker()
        r.visual_k, r.office_shared_pages = 100, 0
        r.visual_seen = r.visual.docs
        r._by_unit = {f"d{i}#p1": [{"chunk_id": f"c{i}", "doc_id": f"d{i}", "format": "pdf", "source": "s",
                                    "pages": [1], "text": "apple"}] for i in range(N)}  # fmt: skip
        allow = frozenset({"d150", "scan"})
        got = [(c["doc_id"], c["chunk_id"]) for c in r.candidates("apple", QVEC, "hybrid_visual_fa", QVEC, allow)]
        self.assertEqual(got, [("d150", "c150"), ("scan", "visual:scan#p1")])


class TestSearchFilter(Tmp):
    def setUp(self):
        super().setUp()
        lance_table(self.cfg)
        update_meta(self.cfg, [DocMeta(f"d{i}", f"d{i}", "xlsx" if i == 150 else "pdf") for i in range(N)])

    def test_search_returns_only_allowed_docs(self):
        from docsearch.search import search

        fake = SimpleNamespace(queries=lambda texts: [QVEC for _ in texts])
        with mock.patch("docsearch.models.Embedder", return_value=fake):
            hits = search(self.cfg, "apple", use_rerank=False, k=5, flt=Filter(formats=("xlsx",)))
            plain = search(self.cfg, "apple", use_rerank=False, k=3)
        self.assertEqual([h["doc_id"] for h in hits], ["d150"])
        self.assertEqual([h["doc_id"] for h in plain], ["d0", "d1", "d2"])

    def test_filter_that_allows_nothing_loads_no_model(self):
        from docsearch.search import search

        with mock.patch("docsearch.models.Embedder", side_effect=AssertionError("model loaded")):
            self.assertEqual(search(self.cfg, "apple", flt=Filter(after="2030-01-01")), [])


class TestCli(Tmp):
    def test_bad_date_is_refused(self):
        from docsearch.cli import main

        err = io.StringIO()
        with mock.patch("sys.stderr", err), self.assertRaises(SystemExit) as e:
            main(["--config", str(self.root / "docsearch.yaml"), "search", "q", "--after", "2024-13-01"])
        self.assertEqual(e.exception.code, 2)
        self.assertIn("YYYY-MM-DD", err.getvalue())

    def test_filter_without_metadata_says_how_to_build_it(self):
        from docsearch.cli import main

        with self.assertRaises(SystemExit) as e:
            main(["--config", str(self.root / "docsearch.yaml"), "search", "q", "--format", "pdf"])
        self.assertIn("docsearch meta", str(e.exception.code))

    def search_cli(self, *args) -> str:
        from docsearch.cli import main

        fake = SimpleNamespace(queries=lambda texts: [QVEC for _ in texts])
        out = io.StringIO()
        with mock.patch("docsearch.models.Embedder", return_value=fake), redirect_stdout(out):
            main(["--config", str(self.root / "docsearch.yaml"), "search", "apple", "--no-rerank", *args])
        return out.getvalue()

    def test_filter_on_an_index_older_than_its_metadata_says_run_meta(self):
        lance_table(self.cfg)
        update_meta(self.cfg, [DocMeta("d150", "d150", "xlsx")])
        with self.assertRaises(SystemExit) as e:
            self.search_cli("--format", "pdf")
        self.assertIn("299 of 300 indexed files have no metadata", str(e.exception.code))
        self.assertIn("uv run docsearch meta", str(e.exception.code))

    def test_filter_with_metadata_for_every_indexed_file_searches(self):
        lance_table(self.cfg)
        update_meta(self.cfg, [DocMeta(f"d{i}", f"d{i}", "xlsx" if i == 150 else "pdf") for i in range(N)])
        self.assertIn("\n    d150\n", self.search_cli("--format", "xlsx"))

    def test_meta_command_writes_rows(self):
        from docsearch.cli import main

        (self.corpus / "a.csv").write_text("x\n")
        write_manifest(self.corpus / "manifest.csv", [Doc("s/a.csv", "a.csv", "csv", "s", "")])
        out = io.StringIO()
        with redirect_stdout(out):
            rc = main(["--config", str(self.root / "docsearch.yaml"), "meta"])
        self.assertEqual(rc, 0)
        row = json.loads(meta_path(self.cfg).read_text())
        self.assertEqual((row["doc_id"], row["file_name"], row["date_source"]), ("s/a.csv", "a.csv", "none"))


if __name__ == "__main__":
    unittest.main()
