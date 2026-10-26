import io
import unittest
from datetime import datetime

from docsearch.cells import docling_grid, grid_cells, is_fused, parse_num, table_quality, xlsx_sheets, xlsx_text


class TestParseNum(unittest.TestCase):
    def test_whole_cell_numbers(self):
        cases = {"(569)": -569.0, "1,320": 1320.0, "4.5%": 4.5, "£44,431": 44431.0, "12.3": 12.3, "-0.5": -0.5,
                 "−1.2": -1.2, "2.6 %": 2.6, "$ 1,234.50": 1234.5, "14 015": 14015.0, "0": 0.0}  # fmt: skip
        self.assertEqual({k: parse_num(k) for k in cases}, cases)

    def test_missing_markers_and_text_are_not_numbers(self):
        for s in (":", "-", "–", "", "Young (15-24)", "78.4 78.5", "339 318", "4.5 pps", "2022-2023", "n/a"):
            self.assertIsNone(parse_num(s), s)

    def test_fused_means_two_numbers_without_a_word(self):
        self.assertEqual([is_fused(s) for s in ("78.4 78.5", "0.1 pps 3.4 %", "1.5 % 5.3 %", "339 318")],
                         [True, True, True, True])  # fmt: skip
        self.assertEqual([is_fused(s) for s in ("Young (15-24)", "12.3", "(569)", ":", "Q1 2023 revenue", "245 W 57TH ST")],
                         [False, False, False, False, False, False])  # fmt: skip
        self.assertTrue(is_fused("7.5 p 6.1 e"))


ANNEX = [["Estonia", "2019", "2022", "2023"],
         ["4 - Activity rate", "78.5", "81.2", "81.6"],
         ["Young (15-24)", "43.5", "44.9", "43.7"],
         ["12 - Unemployment rate", "5.4", "5.6", ":"],
         ["Young (15-24)", "11.1", "(569)", "78.4 78.5"]]  # fmt: skip


def at(cells, row, col):
    return next(c for c in cells if (c["row"], c["col"]) == (row, col))


class TestGridCells(unittest.TestCase):
    def test_labels_headers_sections(self):
        cells = grid_cells(ANNEX)
        self.assertEqual(at(cells, 3, 3), {"row": 3, "col": 3, "text": "44.9", "num": 44.9, "fused": False,
                                           "row_label": "Young (15-24)", "col_header": "2022",
                                           "section": "4 - Activity rate"})  # fmt: skip
        self.assertEqual(at(cells, 5, 3)["num"], -569.0)
        self.assertEqual(at(cells, 5, 3)["section"], "12 - Unemployment rate")
        self.assertEqual((at(cells, 5, 4)["num"], at(cells, 5, 4)["fused"]), (None, True))
        self.assertEqual((at(cells, 4, 4)["num"], at(cells, 4, 4)["fused"]), (None, False))
        self.assertEqual(at(cells, 4, 1)["section"], "12 - Unemployment rate")
        self.assertEqual((at(cells, 1, 2)["col_header"], at(cells, 1, 2)["section"]), (None, None))
        self.assertEqual(len(cells), 20)

    def test_stacked_header_rows_join_and_empty_cells_are_dropped(self):
        grid = [["Region", "2022", "2022"], ["", "Q1", "Q2"], ["North", "5", ""], ["South", "7", "8"]]
        cells = grid_cells(grid)
        self.assertEqual(at(cells, 4, 3)["col_header"], "2022 Q2")
        self.assertEqual(at(cells, 4, 3)["row_label"], "South")
        self.assertEqual(at(cells, 3, 1)["col_header"], "Region")
        self.assertEqual([(c["row"], c["col"]) for c in cells if c["row"] == 3], [(3, 1), (3, 2)])

    def test_text_only_table_takes_its_first_row_as_header(self):
        cells = grid_cells([["Control", "Name"], ["AC-1", "Policy"], ["AC-2", "Account management"]])
        self.assertEqual((at(cells, 3, 2)["col_header"], at(cells, 3, 2)["row_label"]), ("Name", "AC-2"))

    def test_quality_is_filled_share_of_the_grid(self):
        self.assertEqual(table_quality([["a", ""], ["1", "2"]]), (2, 2, 0.75))
        self.assertEqual(table_quality([]), (0, 0, 0.0))


class TestDocling(unittest.TestCase):
    def test_spans_are_expanded(self):
        def cell(r0, r1, c0, c1, text):
            return {"start_row_offset_idx": r0, "end_row_offset_idx": r1, "start_col_offset_idx": c0,
                    "end_col_offset_idx": c1, "text": text}  # fmt: skip

        data = {"num_rows": 2, "num_cols": 3,
                "table_cells": [cell(0, 1, 0, 1, "Country"), cell(0, 1, 1, 3, "2022-2023"),
                                cell(1, 2, 0, 1, "EU"), cell(1, 2, 2, 3, "2.6 %")]}  # fmt: skip
        self.assertEqual(docling_grid(data), [["Country", "2022-2023", "2022-2023"], ["EU", "", "2.6 %"]])


class TestXlsx(unittest.TestCase):
    def test_cell_text_follows_the_number_format(self):
        self.assertEqual(xlsx_text(0.045, "0.0%"), "4.5%")
        self.assertEqual(xlsx_text(1320, "#,##0"), "1320")
        self.assertEqual(xlsx_text(44431.0, "General"), "44431")
        self.assertEqual(xlsx_text(0.1 + 0.2, "General"), "0.3")
        self.assertEqual(xlsx_text(datetime(2020, 1, 6), "yyyy-mm-dd"), "2020-01-06")
        self.assertEqual(xlsx_text(" a ", "General"), "a")
        self.assertEqual(xlsx_text(None, "General"), "")

    def test_hidden_sheets_are_skipped_and_pages_keep_workbook_order(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "README"
        wb.active["A1"] = "about"
        hidden = wb.create_sheet("calc")
        hidden.sheet_state = "hidden"
        hidden["A1"] = 1
        data = wb.create_sheet("Data")
        data.append(["Region", "2022", None])
        data.append([None, "Q1", "Q2"])
        data.append(["North", 5, 6])
        data.merge_cells("B1:C1")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        sheets, skipped = xlsx_sheets(buf)
        self.assertEqual(skipped, ["calc"])
        self.assertEqual([(p, name) for p, name, _ in sheets], [(1, "README"), (3, "Data")])
        self.assertEqual(sheets[1][2], [["Region", "2022", "2022"], ["", "Q1", "Q2"], ["North", "5", "6"]])


if __name__ == "__main__":
    unittest.main()
